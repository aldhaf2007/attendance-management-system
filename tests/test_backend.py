import pytest
import pytest_asyncio
import io
import openpyxl
from datetime import date, time, datetime, timezone, timedelta
from httpx import AsyncClient, ASGITransport
from sqlalchemy.ext.asyncio import create_async_engine, AsyncSession, async_sessionmaker
from sqlalchemy.pool import StaticPool

from app.main import app
from app.database import Base, get_db
from app.models import Department, User, UserRole, Student, AttendanceRecord, AttendanceStatus, Staff, StaffDepartment, CalendarOverride, DayType
from app.security import get_password_hash, get_pin_hash
from app.time_lock import IST, verify_dual_layer_time_lock
from app.config import settings
from app.rate_limiter import rate_limiter

settings.ENVIRONMENT = "test"
settings.ALLOW_TIME_LOCK_BYPASS = True

# In-memory SQLite async engine for tests
TEST_DATABASE_URL = "sqlite+aiosqlite:///:memory:"

test_engine = create_async_engine(
    TEST_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)

TestingSessionLocal = async_sessionmaker(
    bind=test_engine,
    class_=AsyncSession,
    expire_on_commit=False,
    autocommit=False,
    autoflush=False
)

@pytest_asyncio.fixture(scope="session", autouse=True)
async def prepare_database():
    async with test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    yield
    async with test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)

async def override_get_db():
    async with TestingSessionLocal() as session:
        try:
            yield session
            await session.commit()
        except Exception:
            await session.rollback()
            raise

app.dependency_overrides[get_db] = override_get_db

@pytest_asyncio.fixture
async def async_client():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        yield client

@pytest_asyncio.fixture(autouse=True)
async def seed_data():
    rate_limiter.reset_all()
    async with TestingSessionLocal() as session:
        # Clear database tables before seed
        if "calendar_overrides" in Base.metadata.tables:
            await session.execute(Base.metadata.tables["calendar_overrides"].delete())
        await session.execute(Base.metadata.tables["attendance_records"].delete())
        await session.execute(Base.metadata.tables["students"].delete())
        await session.execute(Base.metadata.tables["staffs"].delete())
        if "staff_departments" in Base.metadata.tables:
            await session.execute(Base.metadata.tables["staff_departments"].delete())
        await session.execute(Base.metadata.tables["users"].delete())
        await session.execute(Base.metadata.tables["departments"].delete())
        await session.commit()

        # 1. Create Departments
        dept_cs = Department(id=1, name="Computer Science")
        dept_comm = Department(id=2, name="Commerce")
        session.add_all([dept_cs, dept_comm])
        await session.flush()

        # 2. Create 3-Level Users (Admin, Department Terminal, Staff)
        admin = User(
            id=1,
            username="admin_user",
            password_hash=get_password_hash("admin123"),
            pin_hash=None,
            role=UserRole.ADMIN,
            department_id=None
        )
        dept_terminal = User(
            id=2,
            username="cs_department",
            password_hash=get_password_hash("dept123"),
            pin_hash=None,
            role=UserRole.DEPARTMENT,
            department_id=1
        )
        staff_cs = User(
            id=3,
            username="prof_smith",
            password_hash=get_password_hash("staff123"),
            pin_hash=get_pin_hash("1234"),
            role=UserRole.STAFF,
            department_id=1
        )
        session.add_all([admin, dept_terminal, staff_cs])
        await session.flush()

        # 2.1 Create Staff Department & Staff Model Record
        sd1 = StaffDepartment(id=1, department_1_id=1)
        session.add(sd1)
        await session.flush()

        staff_member = Staff(
            id=1,
            name="Dr. John Smith",
            initials="JS",
            username="prof_smith",
            password_hash=staff_cs.password_hash,
            pin_hash=staff_cs.pin_hash,
            department_id=sd1.id,
            user_id=staff_cs.id
        )
        session.add(staff_member)
        await session.flush()

        # 3. Create Students (CS and Commerce)
        st1 = Student(roll_no="22CS01", name="Alice CS", year=2, department_id=1)
        st2 = Student(roll_no="22CS02", name="Bob CS", year=2, department_id=1)
        st3 = Student(roll_no="22CM01", name="Charlie Commerce", year=2, department_id=2)
        session.add_all([st1, st2, st3])
        await session.flush()

        # 4. Create Attendance Record under Staff JS (id=1)
        today = date.today()
        rec1 = AttendanceRecord(date=today, hour_number=1, roll_no="22CS01", status=AttendanceStatus.PRESENT, staff_id=1)
        rec2 = AttendanceRecord(date=today, hour_number=1, roll_no="22CS02", status=AttendanceStatus.ABSENT, staff_id=1)
        rec3 = AttendanceRecord(date=today, hour_number=3, roll_no="22CM01", status=AttendanceStatus.PRESENT, staff_id=1)
        session.add_all([rec1, rec2, rec3])

        await session.commit()


@pytest.mark.asyncio
async def test_3_level_rbac_roles(async_client: AsyncClient):
    # 1. Level 1 Admin Login & Access
    admin_login = await async_client.post("/api/auth/login", json={"username": "admin_user", "password": "admin123"})
    assert admin_login.status_code == 200
    admin_token = admin_login.json()["access_token"]
    admin_headers = {"Authorization": f"Bearer {admin_token}"}

    # Admin access Admin CRUD -> 200 OK
    admin_depts = await async_client.get("/api/admin/departments", headers=admin_headers)
    assert admin_depts.status_code == 200
    assert len(admin_depts.json()) == 2

    # 2. Level 2 Department Terminal Login & Access
    dept_login = await async_client.post("/api/auth/login", json={"username": "cs_department", "password": "dept123"})
    assert dept_login.status_code == 200
    dept_token = dept_login.json()["access_token"]
    dept_headers = {"Authorization": f"Bearer {dept_token}"}

    # Department Terminal access Admin CRUD -> 403 Forbidden
    dept_admin_res = await async_client.get("/api/admin/departments", headers=dept_headers)
    assert dept_admin_res.status_code == 403

    # Department Terminal access Kiosk Staff List -> 200 OK
    dept_staff_list = await async_client.get("/api/auth/kiosk/staff-list", headers=dept_headers)
    assert dept_staff_list.status_code == 200
    assert len(dept_staff_list.json()) == 1
    assert dept_staff_list.json()[0]["username"] == "prof_smith"

    # 3. Level 3 Staff Login & Access
    staff_login = await async_client.post("/api/auth/login", json={"username": "prof_smith", "password": "staff123"})
    assert staff_login.status_code == 200
    staff_token = staff_login.json()["access_token"]
    staff_headers = {"Authorization": f"Bearer {staff_token}"}

    # Staff access Admin CRUD -> 403 Forbidden
    staff_admin_res = await async_client.get("/api/admin/departments", headers=staff_headers)
    assert staff_admin_res.status_code == 403


@pytest.mark.asyncio
async def test_staff_portal_history(async_client: AsyncClient):
    # Login as Staff prof_smith
    staff_login = await async_client.post("/api/auth/login", json={"username": "prof_smith", "password": "staff123"})
    staff_token = staff_login.json()["access_token"]
    staff_headers = {"Authorization": f"Bearer {staff_token}"}

    # 1. GET /api/staff/my-history (Combined across all departments)
    res = await async_client.get("/api/staff/my-history", headers=staff_headers)
    assert res.status_code == 200
    data = res.json()
    
    assert data["staff_username"] == "prof_smith"
    assert data["staff_name"] == "Dr. John Smith"
    assert data["primary_department_name"] == "Computer Science"
    assert data["total_classes_conducted"] == 2
    assert data["total_records_logged"] == 3
    assert len(data["records"]) == 3
    assert len(data["departments"]) == 2

    # Check department summaries
    dept_names = [d["department_name"] for d in data["departments"]]
    assert "Computer Science" in dept_names
    assert "Commerce" in dept_names

    # 2. Filter by Commerce (department_id=2)
    comm_res = await async_client.get("/api/staff/my-history?department_id=2", headers=staff_headers)
    assert comm_res.status_code == 200
    comm_data = comm_res.json()
    assert comm_data["selected_department_id"] == 2
    assert comm_data["total_classes_conducted"] == 1
    assert comm_data["total_records_logged"] == 1
    assert len(comm_data["records"]) == 1
    assert comm_data["records"][0]["department_name"] == "Commerce"
    assert comm_data["records"][0]["roll_no"] == "22CM01"

    # 3. Filter by CS (department_id=1)
    cs_res = await async_client.get("/api/staff/my-history?department_id=1", headers=staff_headers)
    assert cs_res.status_code == 200
    cs_data = cs_res.json()
    assert cs_data["selected_department_id"] == 1
    assert cs_data["total_records_logged"] == 2


@pytest.mark.asyncio
async def test_academic_calendar_and_overrides(async_client: AsyncClient):
    # 1. Login as Admin
    admin_login = await async_client.post("/api/auth/login", json={"username": "admin_user", "password": "admin123"})
    assert admin_login.status_code == 200
    admin_headers = {"Authorization": f"Bearer {admin_login.json()['access_token']}"}

    # 2. Create a Holiday Override
    holiday_date = (date.today() + timedelta(days=1)).isoformat()
    create_res = await async_client.post(
        "/api/admin/calendar/override",
        headers=admin_headers,
        json={
            "date": holiday_date,
            "day_type": "holiday",
            "active_periods": [],
            "description": "Founder's Day Holiday"
        }
    )
    assert create_res.status_code == 200
    assert create_res.json()["day_type"] == "holiday"

    # 3. Verify Override appears in Admin Overrides List
    list_res = await async_client.get("/api/admin/calendar/overrides", headers=admin_headers)
    assert list_res.status_code == 200
    overrides = list_res.json()
    assert any(o["date"] == holiday_date and o["day_type"] == "holiday" for o in overrides)

    # 4. Login as Staff (prof_smith) who has authorization to submit attendance
    staff_login = await async_client.post("/api/auth/login", json={"username": "prof_smith", "password": "staff123"})
    assert staff_login.status_code == 200
    staff_headers = {
        "Authorization": f"Bearer {staff_login.json()['access_token']}",
        "X-Bypass-Time-Lock": "bypass-secret-test"
    }

    # 5. Attempt attendance submission on Holiday Date -> Expect 403 Forbidden
    sub_holiday = await async_client.post(
        "/api/attendance/submit",
        headers=staff_headers,
        json={
            "date": holiday_date,
            "hour_number": 1,
            "records": [{"roll_no": "22CS01", "status": "Present"}]
        }
    )
    assert sub_holiday.status_code == 403
    assert "holiday" in sub_holiday.json()["detail"].lower()

    # 6. Create a Half-Day Override (Periods 1-2 active)
    half_date = (date.today() + timedelta(days=2)).isoformat()
    half_res = await async_client.post(
        "/api/admin/calendar/override",
        headers=admin_headers,
        json={
            "date": half_date,
            "day_type": "half_day",
            "active_periods": [1, 2],
            "description": "Half Day Symposium"
        }
    )
    assert half_res.status_code == 200

    # 7. Attempt submitting an inactive period (Period 4) on Half-Day -> Expect 403 Forbidden
    sub_inactive = await async_client.post(
        "/api/attendance/submit",
        headers=staff_headers,
        json={
            "date": half_date,
            "hour_number": 4,
            "records": [{"roll_no": "22CS01", "status": "Present"}]
        }
    )
    assert sub_inactive.status_code == 403
    assert "not an active period" in sub_inactive.json()["detail"].lower()

    # 8. Submitting an active period (Period 1) on Half-Day with bypass -> 200 OK
    sub_active = await async_client.post(
        "/api/attendance/submit",
        headers=staff_headers,
        json={
            "date": half_date,
            "hour_number": 1,
            "records": [{"roll_no": "22CS01", "status": "Present"}]
        }
    )
    assert sub_active.status_code == 200

    # 9. Verify Calendar status in Attendance Grid response
    grid_res = await async_client.get(
        f"/api/attendance/grid?department_id=1&year=2&target_date={half_date}",
        headers=staff_headers
    )
    assert grid_res.status_code == 200
    grid_data = grid_res.json()
    assert grid_data["day_type"] == "half_day"
    assert grid_data["active_periods"] == [1, 2]
    assert grid_data["day_description"] == "Half Day Symposium"

    # 10. Admin Delete Override
    del_res = await async_client.delete(f"/api/admin/calendar/override/{holiday_date}", headers=admin_headers)
    assert del_res.status_code == 200


@pytest.mark.asyncio
async def test_export_endpoints(async_client: AsyncClient):
    # Login as Admin
    admin_login = await async_client.post("/api/auth/login", json={"username": "admin_user", "password": "admin123"})
    assert admin_login.status_code == 200
    admin_headers = {"Authorization": f"Bearer {admin_login.json()['access_token']}"}

    # 1. Test Available Years
    years_res = await async_client.get("/api/export/available-years", headers=admin_headers)
    assert years_res.status_code == 200
    years_data = years_res.json()
    assert "years" in years_data
    assert isinstance(years_data["years"], list)
    assert date.today().year in years_data["years"]

    # 2. Test Full Academic Year Export
    yearly_res = await async_client.get(
        f"/api/export/attendance/1/2/1?export_type=yearly&year_date={date.today().year}",
        headers=admin_headers
    )
    assert yearly_res.status_code == 200
    assert yearly_res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "FullYear" in yearly_res.headers["content-disposition"]

    # 3. Test Custom Date Range Export
    start = f"{date.today().year}-01-01"
    end = f"{date.today().year}-12-31"
    custom_res = await async_client.get(
        f"/api/export/attendance/1/2/1?export_type=custom&start_date={start}&end_date={end}",
        headers=admin_headers
    )
    assert custom_res.status_code == 200
    assert custom_res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "Range" in custom_res.headers["content-disposition"]


@pytest.mark.asyncio
async def test_kiosk_direct_unlock(async_client: AsyncClient):
    # 1. Direct Initials + PIN unlock (Initials "JS" + PIN "1234" belongs to Dr. John Smith)
    init_res = await async_client.post("/api/auth/kiosk/direct-unlock", json={"initials": "JS", "pin": "1234", "hour_number": 1})
    assert init_res.status_code == 200
    init_data = init_res.json()
    assert init_data["staff_name"] == "Dr. John Smith"
    assert init_data["staff_initials"] == "JS"
    assert init_data["staff_username"] == "prof_smith"
    assert init_data["department_id"] == 1
    assert "access_token" in init_data

    # 2. Case insensitive initials ("js")
    case_res = await async_client.post("/api/auth/kiosk/direct-unlock", json={"initials": "js", "pin": "1234", "hour_number": 1})
    assert case_res.status_code == 200
    assert case_res.json()["staff_initials"] == "JS"

    # 3. Direct PIN-only unlock fallback
    pin_res = await async_client.post("/api/auth/kiosk/direct-unlock", json={"pin": "1234", "hour_number": 1})
    assert pin_res.status_code == 200
    assert pin_res.json()["staff_initials"] == "JS"

    # 4. Invalid initials -> 404
    invalid_init_res = await async_client.post("/api/auth/kiosk/direct-unlock", json={"initials": "XX", "pin": "1234", "hour_number": 1})
    assert invalid_init_res.status_code == 404

    # 5. Valid initials but wrong PIN -> 401
    invalid_pin_res = await async_client.post("/api/auth/kiosk/direct-unlock", json={"initials": "JS", "pin": "9999", "hour_number": 1})
    assert invalid_pin_res.status_code == 401

    # 6. Verify Grid returns period_staff with staff details using the kiosk access_token
    token = init_data["access_token"]
    grid_res = await async_client.get(
        f"/api/attendance/grid?department_id=1&year=2&target_date={date.today().isoformat()}",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert grid_res.status_code == 200
    grid_data = grid_res.json()
    assert "period_staff" in grid_data
    # Hour 1 was recorded by Staff 1 (JS)
    p1_staff = grid_data["period_staff"].get("1") or grid_data["period_staff"].get(1)
    assert p1_staff is not None
    assert p1_staff["initials"] == "JS"
    assert p1_staff["name"] == "Dr. John Smith"


@pytest.mark.asyncio
async def test_staff_excel_export_and_filtering(async_client: AsyncClient):
    import io
    import openpyxl

    # 1. Login as Staff prof_smith
    staff_login = await async_client.post("/api/auth/login", json={"username": "prof_smith", "password": "staff123"})
    assert staff_login.status_code == 200
    staff_token = staff_login.json()["access_token"]
    staff_headers = {"Authorization": f"Bearer {staff_token}"}

    # 2. Test Timeframe Filtering: All Days vs Yearly vs Monthly
    # All Days
    all_res = await async_client.get("/api/staff/my-history?filter_type=all", headers=staff_headers)
    assert all_res.status_code == 200
    all_data = all_res.json()
    assert all_data["filter_type"] == "all"
    assert len(all_data["records"]) == 3
    assert all_data["available_years"] is not None

    # Yearly filter for current year
    curr_year = date.today().year
    year_res = await async_client.get(f"/api/staff/my-history?filter_type=yearly&year_date={curr_year}", headers=staff_headers)
    assert year_res.status_code == 200
    year_data = year_res.json()
    assert year_data["filter_type"] == "yearly"
    assert all(r["date"].startswith(str(curr_year)) for r in year_data["records"])

    # Monthly filter
    curr_month = date.today().month
    month_res = await async_client.get(f"/api/staff/my-history?filter_type=monthly&year_date={curr_year}&month={curr_month}", headers=staff_headers)
    assert month_res.status_code == 200

    # 3. Test Staff Excel Export: All records
    export_all = await async_client.get("/api/staff/export-excel?export_type=all", headers=staff_headers)
    assert export_all.status_code == 200
    assert export_all.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "Attendance_Staff_JS" in export_all.headers["content-disposition"]

    # Verify Excel workbook content with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(export_all.content))
    assert "Staff Teaching Summary" in wb.sheetnames
    assert "Detailed Periods Log" in wb.sheetnames

    # Check Sheet 1 (Summary Roster & KPIs)
    ws1 = wb["Staff Teaching Summary"]
    assert "FACULTY ATTENDANCE REPORT" in ws1.cell(row=2, column=2).value
    assert "Dr. John Smith" in ws1.cell(row=3, column=2).value or "JS" in ws1.cell(row=3, column=2).value

    # Check Sheet 2 (Detailed Periods Log)
    ws2 = wb["Detailed Periods Log"]
    # Header row is row 5
    assert ws2.cell(row=5, column=2).value == "Date"
    assert ws2.cell(row=5, column=4).value == "Period"
    # All data rows in Sheet 2 must be for Dr. John Smith / JS
    row_idx = 6
    while ws2.cell(row=row_idx, column=2).value is not None:
        faculty_name = ws2.cell(row=row_idx, column=10).value
        faculty_init = ws2.cell(row=row_idx, column=11).value
        assert faculty_name == "Dr. John Smith"
        assert faculty_init == "JS"
        row_idx += 1
    assert row_idx > 6  # Verified at least one record exists

    # 4. Test Filtered Staff Excel Export: Department-scoped
    export_dept = await async_client.get("/api/staff/export-excel?department_id=1&export_type=all", headers=staff_headers)
    assert export_dept.status_code == 200
    wb_dept = openpyxl.load_workbook(io.BytesIO(export_dept.content))
    ws2_dept = wb_dept["Detailed Periods Log"]
    r_idx = 6
    while ws2_dept.cell(row=r_idx, column=2).value is not None:
        dept_name = ws2_dept.cell(row=r_idx, column=5).value
        assert dept_name == "Computer Science"
        r_idx += 1


@pytest.mark.asyncio
async def test_create_department_without_password(async_client: AsyncClient):
    # 1. Login as Admin
    admin_login = await async_client.post("/api/auth/login", json={"username": "admin_user", "password": "admin123"})
    assert admin_login.status_code == 200
    admin_headers = {"Authorization": f"Bearer {admin_login.json()['access_token']}"}

    # 2. Create department with name only (no terminal credentials)
    create_res1 = await async_client.post(
        "/api/admin/departments",
        headers=admin_headers,
        json={"name": "Civil Engineering"}
    )
    assert create_res1.status_code == 200
    dept1 = create_res1.json()
    assert dept1["name"] == "Civil Engineering"
    assert "id" in dept1
    # Check that a dashboard username was auto-generated
    assert dept1.get("account_username") == "civil_engineering_dept"

    # Verify auto-generated dashboard account can log in immediately
    dept1_login = await async_client.post(
        "/api/auth/login",
        json={"username": "civil_engineering_dept", "password": ""}
    )
    assert dept1_login.status_code == 200
    assert dept1_login.json()["role"] == "Department"
    assert dept1_login.json()["department_id"] == dept1["id"]

    # Verify dashboard user has access to Department Dashboard APIs
    dept1_headers = {"Authorization": f"Bearer {dept1_login.json()['access_token']}"}
    dash_res = await async_client.get(f"/api/analytics/attendance-summary?department_id={dept1['id']}&year=2", headers=dept1_headers)
    assert dash_res.status_code == 200

    # 3. Create department with name and terminal username, but NO password
    create_res2 = await async_client.post(
        "/api/admin/departments",
        headers=admin_headers,
        json={"name": "Mechanical Engineering", "account_username": "mech_department"}
    )
    assert create_res2.status_code == 200
    dept2 = create_res2.json()
    assert dept2["name"] == "Mechanical Engineering"
    assert "id" in dept2

    # 4. Verify terminal user created without password can authenticate
    dept_login = await async_client.post(
        "/api/auth/login",
        json={"username": "mech_department", "password": ""}
    )
    assert dept_login.status_code == 200
    assert dept_login.json()["role"] == "Department"
    assert dept_login.json()["department_id"] == dept2["id"]

@pytest.mark.anyio
async def test_create_and_update_department_dashboard_password(async_client: AsyncClient):
    # 1. Admin login
    admin_login = await async_client.post("/api/auth/login", json={"username": "admin_user", "password": "admin123"})
    assert admin_login.status_code == 200
    admin_token = admin_login.json()["access_token"]
    admin_headers = {"Authorization": f"Bearer {admin_token}"}

    # 2. Create department with username and password
    create_res = await async_client.post(
        "/api/admin/departments",
        headers=admin_headers,
        json={
            "name": "BioTechnology",
            "account_username": "biotech_dept",
            "account_password": "biotechsecret"
        }
    )
    assert create_res.status_code == 200
    dept = create_res.json()
    assert dept["name"] == "BioTechnology"
    assert dept["account_username"] == "biotech_dept"

    # 3. Try login with incorrect password -> 401
    bad_login = await async_client.post(
        "/api/auth/login",
        json={"username": "biotech_dept", "password": "wrongpassword"}
    )
    assert bad_login.status_code == 401

    # 4. Login with correct password -> 200 OK, role Department
    good_login = await async_client.post(
        "/api/auth/login",
        json={"username": "biotech_dept", "password": "biotechsecret"}
    )
    assert good_login.status_code == 200
    assert good_login.json()["role"] == "Department"
    assert good_login.json()["department_id"] == dept["id"]

    # 5. Update department password
    update_res = await async_client.put(
        f"/api/admin/departments/{dept['id']}",
        headers=admin_headers,
        json={
            "account_password": "newbiotechsecret"
        }
    )
    assert update_res.status_code == 200

    # Old password no longer works
    old_login = await async_client.post(
        "/api/auth/login",
        json={"username": "biotech_dept", "password": "biotechsecret"}
    )
    assert old_login.status_code == 401

    # New password works
    new_login = await async_client.post(
        "/api/auth/login",
        json={"username": "biotech_dept", "password": "newbiotechsecret"}
    )
    assert new_login.status_code == 200
    assert new_login.json()["role"] == "Department"

@pytest.mark.asyncio
async def test_kiosk_multi_department_selection(async_client: AsyncClient):
    # 1. Admin login to configure additional departments
    admin_login = await async_client.post("/api/auth/login", json={"username": "admin_user", "password": "admin123"})
    assert admin_login.status_code == 200
    admin_token = admin_login.json()["access_token"]
    admin_headers = {"Authorization": f"Bearer {admin_token}"}

    # 2. Update staff member 1 (Dr. John Smith, initials 'JS') to also teach in Department 2 (Commerce)
    update_res = await async_client.put(
        "/api/admin/staff/1",
        headers=admin_headers,
        json={
            "name": "Dr. John Smith",
            "initials": "JS",
            "username": "prof_smith",
            "department_id": 1,
            "additional_department_ids": [2]
        }
    )
    assert update_res.status_code == 200
    staff_data = update_res.json()
    assert staff_data["additional_department_ids"] == [2]

    # Verify admin staff list returns additional departments
    staff_list_res = await async_client.get("/api/admin/staff", headers=admin_headers)
    assert staff_list_res.status_code == 200
    staff_members = staff_list_res.json()
    smith = next(s for s in staff_members if s["id"] == 1)
    assert smith["additional_department_ids"] == [2]
    assert "Commerce" in smith.get("additional_departments", [])

    # 3. Verify public kiosk departments list endpoint
    kiosk_depts_res = await async_client.get("/api/auth/kiosk/departments")
    assert kiosk_depts_res.status_code == 200
    kiosk_depts = kiosk_depts_res.json()
    assert len(kiosk_depts) >= 2
    dept_ids = [d["id"] for d in kiosk_depts]
    assert 1 in dept_ids and 2 in dept_ids

    # 4. Direct unlock at Attendance Terminal with Staff Initials 'JS' and PIN '1234'
    unlock_res = await async_client.post(
        "/api/auth/kiosk/direct-unlock",
        json={
            "initials": "JS",
            "pin": "1234",
            "hour_number": 1
        }
    )
    assert unlock_res.status_code == 200
    unlock_data = unlock_res.json()
    assert unlock_data["staff_initials"] == "JS"
    assert "departments" in unlock_data
    # Must offer both Computer Science (1) and Commerce (2)
    offered_dept_ids = [d["id"] for d in unlock_data["departments"]]
    assert 1 in offered_dept_ids
    assert 2 in offered_dept_ids

    # Verify primary vs cross-teaching flags
    primary_opt = next(d for d in unlock_data["departments"] if d["id"] == 1)
    assert primary_opt["is_primary"] is True
    secondary_opt = next(d for d in unlock_data["departments"] if d["id"] == 2)
    assert secondary_opt["is_primary"] is False

    # 5. Staff chooses Department 2 (Commerce) at terminal
    select_res = await async_client.post(
        "/api/auth/kiosk/select-department",
        headers={"Authorization": f"Bearer {unlock_data['access_token']}"},
        json={
            "staff_id": 1,
            "department_id": 2,
            "hour_number": 1
        }
    )
    assert select_res.status_code == 200
    select_data = select_res.json()
    assert select_data["department_id"] == 2
    assert select_data["department_name"] == "Commerce"
    commerce_token = select_data["access_token"]

    # 6. Staff takes attendance for Commerce Class (Dept 2, Year 2) using this token
    # (Sep 7, 2026 is Monday, an active academic day)
    academic_date = "2026-09-07"
    commerce_sub_res = await async_client.post(
        "/api/attendance/submit",
        headers={
            "Authorization": f"Bearer {commerce_token}",
            "X-Bypass-Time-Lock": "bypass-secret-test"
        },
        json={
            "date": academic_date,
            "hour_number": 1,
            "records": [
                {"roll_no": "22CM01", "status": "Present"}
            ]
        }
    )
    assert commerce_sub_res.status_code == 200
    sub_records = commerce_sub_res.json()
    assert len(sub_records) == 1
    assert sub_records[0]["roll_no"] == "22CM01"
    assert sub_records[0]["status"] == "Present"


@pytest.mark.asyncio
async def test_security_backdoors_eliminated(async_client: AsyncClient):
    # Verify former hardcoded backdoor credentials fail with 401
    res1 = await async_client.post("/api/auth/login", json={"username": "admin", "password": "admin"})
    assert res1.status_code == 401

    res2 = await async_client.post("/api/auth/login", json={"username": "cs_department", "password": "dept"})
    assert res2.status_code == 401

    res3 = await async_client.post("/api/auth/login", json={"username": "prof_smith", "password": "staff"})
    assert res3.status_code == 401


@pytest.mark.asyncio
async def test_security_headers_present(async_client: AsyncClient):
    # Verify HTTP security headers on responses
    res = await async_client.get("/")
    assert res.status_code == 200
    assert res.headers.get("X-Content-Type-Options") == "nosniff"
    assert res.headers.get("X-Frame-Options") == "DENY"
    assert res.headers.get("X-XSS-Protection") == "1; mode=block"
    assert res.headers.get("Referrer-Policy") == "strict-origin-when-cross-origin"


@pytest.mark.asyncio
async def test_security_pin_brute_force_lockout(async_client: AsyncClient):
    # Staff 'JS' has valid PIN '1234'
    # Test that 5 consecutive wrong PINs trigger account lockout (429)
    for i in range(4):
        wrong_res = await async_client.post(
            "/api/auth/kiosk/direct-unlock",
            json={"initials": "JS", "pin": "0000", "hour_number": 1}
        )
        assert wrong_res.status_code == 401
        assert "attempts remaining" in wrong_res.json()["detail"]

    # 5th failed attempt -> locks account with 429
    fifth_res = await async_client.post(
        "/api/auth/kiosk/direct-unlock",
        json={"initials": "JS", "pin": "0000", "hour_number": 1}
    )
    assert fifth_res.status_code == 429
    assert "locked for 15 minutes" in fifth_res.json()["detail"].lower()

    # 6th attempt (even with the CORRECT PIN) is rejected with 429 because account is locked!
    locked_res = await async_client.post(
        "/api/auth/kiosk/direct-unlock",
        json={"initials": "JS", "pin": "1234", "hour_number": 1}
    )
    assert locked_res.status_code == 429
    assert "locked" in locked_res.json()["detail"].lower()


@pytest.mark.asyncio
async def test_security_token_revocation_on_logout(async_client: AsyncClient):
    # 1. Login as Admin
    login_res = await async_client.post("/api/auth/login", json={"username": "admin_user", "password": "admin123"})
    assert login_res.status_code == 200
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 2. Access protected endpoint -> 200
    me_res = await async_client.get("/api/auth/me", headers=headers)
    assert me_res.status_code == 200
    assert me_res.json()["username"] == "admin_user"

    # 3. Call /api/auth/logout to invalidate token
    logout_res = await async_client.post("/api/auth/logout", headers=headers)
    assert logout_res.status_code == 200

    # 4. Attempt to use revoked token -> 401 Unauthorized
    revoked_res = await async_client.get("/api/auth/me", headers=headers)
    assert revoked_res.status_code == 401
    assert "revoked" in revoked_res.json()["detail"].lower()


@pytest.mark.asyncio
async def test_security_excel_formula_injection_sanitized(async_client: AsyncClient):
    # 1. Login as Admin
    admin_login = await async_client.post("/api/auth/login", json={"username": "admin_user", "password": "admin123"})
    assert admin_login.status_code == 200
    admin_headers = {"Authorization": f"Bearer {admin_login.json()['access_token']}"}

    # 2. Create student with formula injection payload as name
    payload_name = "=CMD|' /C calc'!A0"
    create_res = await async_client.post(
        "/api/admin/students",
        headers=admin_headers,
        json={"roll_no": "22CS99", "name": payload_name, "year": 2, "department_id": 1}
    )
    assert create_res.status_code == 200

    # 3. Export Excel
    export_res = await async_client.get(
        f"/api/export/attendance/1/2/1?export_type=yearly&year_date={date.today().year}",
        headers=admin_headers
    )
    assert export_res.status_code == 200

    # 4. Load workbook and verify formula injection is neutralized with leading quote
    wb = openpyxl.load_workbook(io.BytesIO(export_res.content))
    ws = wb["Class Summary"]
    found_student = False
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and "calc" in str(cell.value):
                found_student = True
                assert str(cell.value).startswith("'="), f"Cell value '{cell.value}' was not sanitized against formula injection!"
    assert found_student, "Student with formula payload was not found in exported Excel sheet"






