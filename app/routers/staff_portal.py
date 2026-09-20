import io
import calendar
from datetime import date
from typing import List, Optional, Dict
from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_, extract
from sqlalchemy.orm import selectinload

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database import get_db
from app.models import AttendanceRecord, Student, Department, AttendanceStatus, Staff, StaffDepartment
from app.schemas import StaffHistoryRecord, StaffHistoryResponse, StaffDepartmentSummary
from app.dependencies import CurrentUserContext, require_staff

router = APIRouter(prefix="/staff", tags=["Staff Personal Portal"])

@router.get("/my-history", response_model=StaffHistoryResponse)
async def get_staff_personal_history(
    department_id: Optional[int] = None,
    filter_type: Optional[str] = Query("all", description="all, yearly, monthly, daily, custom"),
    year_date: Optional[int] = Query(None),
    month: Optional[int] = Query(None),
    target_date: Optional[str] = Query(None),
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    ctx: CurrentUserContext = Depends(require_staff),
    db: AsyncSession = Depends(get_db)
):
    """
    Level 3 Staff Personal Portal Endpoint:
    Fetches read-only attendance history for the logged-in staff member.
    Arranged for staff members who teach across multiple departments:
    - Provides a per-department summary breakdown (classes taught, records, attendance %).
    - Supports filtering across all days, specific year, specific month, specific day, or custom range.
    - Staff members cannot modify past records from this portal.
    """
    staff_id = ctx.staff_id or ctx.user_id

    # 1. Fetch Staff profile details
    staff_res = await db.execute(
        select(Staff).options(selectinload(Staff.staff_department).selectinload(StaffDepartment.dept_1)).where(Staff.id == staff_id)
    )
    staff_obj = staff_res.scalar_one_or_none()

    staff_name = staff_obj.name if staff_obj else None
    staff_initials = staff_obj.initials if staff_obj else None
    primary_dept_name = staff_obj.department.name if staff_obj and staff_obj.department else None

    # 2. Query all attendance records logged by this staff_id
    base_query = (
        select(AttendanceRecord, Student, Department)
        .join(Student, AttendanceRecord.roll_no == Student.roll_no)
        .join(Department, Student.department_id == Department.id)
        .where(AttendanceRecord.staff_id == staff_id)
        .order_by(AttendanceRecord.date.desc(), AttendanceRecord.hour_number.asc(), Student.roll_no.asc())
    )

    result = await db.execute(base_query)
    all_rows = result.all()

    # Available years from distinct record dates
    raw_years = {row[0].date.year for row in all_rows if row[0].date}
    current_year = date.today().year
    raw_years.add(current_year)
    available_years = sorted(list(raw_years), reverse=True)

    # 3. Aggregate per-department analytics across all records
    dept_stats: Dict[int, dict] = {}

    for rec, student, dept in all_rows:
        d_id = dept.id
        if d_id not in dept_stats:
            dept_stats[d_id] = {
                "department_id": d_id,
                "department_name": dept.name,
                "unique_classes": set(),
                "records_count": 0,
                "present_count": 0,
                "absent_count": 0,
                "od_count": 0,
            }

        st = dept_stats[d_id]
        st["unique_classes"].add((rec.date, rec.hour_number, student.department_id, student.year))
        st["records_count"] += 1

        if rec.status == AttendanceStatus.PRESENT:
            st["present_count"] += 1
        elif rec.status == AttendanceStatus.ABSENT:
            st["absent_count"] += 1
        elif rec.status == AttendanceStatus.OD:
            st["od_count"] += 1
            st["present_count"] += 1

    department_summaries: List[StaffDepartmentSummary] = []
    for d_id, st in sorted(dept_stats.items(), key=lambda x: x[1]["department_name"]):
        total_recs = st["records_count"]
        pct = round((st["present_count"] / total_recs) * 100.0, 1) if total_recs > 0 else 100.0
        department_summaries.append(
            StaffDepartmentSummary(
                department_id=st["department_id"],
                department_name=st["department_name"],
                classes_conducted=len(st["unique_classes"]),
                records_logged=total_recs,
                attendance_percentage=pct,
                present_count=st["present_count"],
                absent_count=st["absent_count"],
                od_count=st["od_count"]
            )
        )

    # 4. Filter records by department_id and timeframe
    filtered_rows = all_rows
    if department_id is not None:
        filtered_rows = [row for row in filtered_rows if row[1].department_id == department_id]

    filter_label = "All Days & Years"
    if filter_type == "yearly" and year_date:
        filtered_rows = [row for row in filtered_rows if row[0].date.year == year_date]
        filter_label = f"Year {year_date}"
    elif filter_type == "monthly" and year_date and month:
        filtered_rows = [row for row in filtered_rows if row[0].date.year == year_date and row[0].date.month == month]
        filter_label = f"{calendar.month_name[month]} {year_date}"
    elif filter_type == "daily" and target_date:
        try:
            parsed_d = date.fromisoformat(target_date)
            filtered_rows = [row for row in filtered_rows if row[0].date == parsed_d]
            filter_label = f"Day: {parsed_d.strftime('%b %d, %Y')}"
        except ValueError:
            pass
    elif filter_type == "custom" and start_date and end_date:
        try:
            parsed_start = date.fromisoformat(start_date)
            parsed_end = date.fromisoformat(end_date)
            filtered_rows = [row for row in filtered_rows if parsed_start <= row[0].date <= parsed_end]
            filter_label = f"{parsed_start.strftime('%b %d, %Y')} – {parsed_end.strftime('%b %d, %Y')}"
        except ValueError:
            pass

    history_records: List[StaffHistoryRecord] = []
    active_present = 0
    active_unique_classes = set()

    for rec, student, dept in filtered_rows:
        active_unique_classes.add((rec.date, rec.hour_number, student.department_id, student.year))
        if rec.status == AttendanceStatus.PRESENT or rec.status == AttendanceStatus.OD:
            active_present += 1

        history_records.append(
            StaffHistoryRecord(
                id=rec.id,
                date=rec.date,
                hour_number=rec.hour_number,
                roll_no=student.roll_no,
                student_name=student.name,
                status=rec.status,
                department_id=dept.id,
                department_name=dept.name,
                year=student.year
            )
        )

    total_filtered = len(filtered_rows)
    avg_percentage = round((active_present / total_filtered) * 100.0, 1) if total_filtered > 0 else 100.0

    return StaffHistoryResponse(
        staff_id=staff_id,
        staff_username=ctx.username,
        staff_name=staff_name,
        staff_initials=staff_initials,
        primary_department_name=primary_dept_name,
        selected_department_id=department_id,
        total_classes_conducted=len(active_unique_classes),
        total_records_logged=total_filtered,
        average_attendance_percentage=avg_percentage,
        filter_type=filter_type or "all",
        filter_label=filter_label,
        available_years=available_years,
        departments=department_summaries,
        records=history_records
    )


@router.get("/export-excel")
async def export_staff_attendance_excel(
    department_id: Optional[int] = Query(None),
    export_type: Optional[str] = Query("all", description="all, yearly, monthly, custom"),
    year_date: Optional[int] = Query(2026),
    month: Optional[int] = Query(1),
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    ctx: CurrentUserContext = Depends(require_staff),
    db: AsyncSession = Depends(get_db)
):
    """
    Dedicated Staff Excel Export Endpoint:
    Exports an executive .xlsx workbook containing ONLY the periods and attendance records
    conducted/taken by the logged-in staff member.
    """
    staff_id = ctx.staff_id or ctx.user_id

    # 1. Fetch Staff profile
    staff_res = await db.execute(
        select(Staff).options(selectinload(Staff.staff_department).selectinload(StaffDepartment.dept_1)).where(Staff.id == staff_id)
    )
    staff_obj = staff_res.scalar_one_or_none()

    staff_name = staff_obj.name if staff_obj else ctx.username
    staff_initials = staff_obj.initials if staff_obj else "STAFF"
    primary_dept = staff_obj.department.name if staff_obj and staff_obj.department else "General"

    # 2. Scope Department
    dept_label = "All Departments"
    if department_id is not None:
        d_res = await db.execute(select(Department).where(Department.id == department_id))
        d_obj = d_res.scalar_one_or_none()
        if d_obj:
            dept_label = d_obj.name

    # 3. Base Conditions: EXCLUSIVELY records taken by this staff member
    conditions = [AttendanceRecord.staff_id == staff_id]
    if department_id is not None:
        conditions.append(Student.department_id == department_id)

    date_label = "All Recorded Periods (All Time)"
    filename_suffix = "All_Periods"

    if export_type == "yearly" and year_date:
        conditions.append(extract('year', AttendanceRecord.date) == year_date)
        date_label = f"Academic Year {year_date}"
        filename_suffix = f"FullYear_{year_date}"
    elif export_type == "monthly" and year_date and month:
        if month < 1 or month > 12:
            month = 1
        month_name = calendar.month_name[month]
        conditions.append(extract('year', AttendanceRecord.date) == year_date)
        conditions.append(extract('month', AttendanceRecord.date) == month)
        date_label = f"{month_name} {year_date}"
        filename_suffix = f"{month_name[:3]}_{year_date}"
    elif export_type == "custom" and start_date and end_date:
        try:
            parsed_start = date.fromisoformat(start_date)
            parsed_end = date.fromisoformat(end_date)
            conditions.append(AttendanceRecord.date >= parsed_start)
            conditions.append(AttendanceRecord.date <= parsed_end)
            date_label = f"{parsed_start} to {parsed_end}"
            filename_suffix = f"Range_{parsed_start}_to_{parsed_end}"
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid date format. Use YYYY-MM-DD.")

    query = (
        select(AttendanceRecord, Student, Department)
        .join(Student, AttendanceRecord.roll_no == Student.roll_no)
        .join(Department, Student.department_id == Department.id)
        .where(and_(*conditions))
        .order_by(AttendanceRecord.date.desc(), AttendanceRecord.hour_number.asc(), Student.roll_no.asc())
    )
    res = await db.execute(query)
    all_recs = res.all()

    # Aggregate student stats across records taken by this staff
    student_stats: Dict[str, dict] = {}
    unique_sessions = set()

    for rec, st, dept in all_recs:
        unique_sessions.add((rec.date, rec.hour_number, dept.id, st.year))
        if st.roll_no not in student_stats:
            student_stats[st.roll_no] = {
                "roll_no": st.roll_no,
                "name": st.name,
                "department": dept.name,
                "year": st.year,
                "total_conducted": 0,
                "present": 0,
                "absent": 0,
                "od": 0
            }
        s = student_stats[st.roll_no]
        s["total_conducted"] += 1
        if rec.status == AttendanceStatus.PRESENT:
            s["present"] += 1
        elif rec.status == AttendanceStatus.ABSENT:
            s["absent"] += 1
        elif rec.status == AttendanceStatus.OD:
            s["od"] += 1

    wb = openpyxl.Workbook()

    # ----------------------------------------------------
    # SHEET 1: CLASS / ROSTER SUMMARY
    # ----------------------------------------------------
    ws1 = wb.active
    ws1.title = "Staff Teaching Summary"
    ws1.views.sheetView[0].showGridLines = True

    # Styling
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    green_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    red_fill = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Title & Metadata
    ws1.cell(row=2, column=2, value="FACULTY ATTENDANCE REPORT — PERIODS CONDUCTED BY STAFF").font = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
    ws1.cell(row=3, column=2, value=f"Faculty: {staff_name} ({staff_initials}) | Primary Department: {primary_dept}").font = Font(name="Segoe UI", size=10, bold=True, color="334155")
    ws1.cell(row=4, column=2, value=f"Scope: {dept_label} | Timeframe: {date_label}").font = Font(name="Segoe UI", size=10, color="475569")
    ws1.cell(row=5, column=2, value=f"* Scope Constraint: This report exclusively includes attendance records for periods conducted by {staff_name}.").font = Font(name="Segoe UI", size=9, italic=True, color="6D28D9")

    # KPI Block (Row 7-8)
    kpis = [
        ("Students Taught", len(student_stats)),
        ("Classes Conducted", len(unique_sessions)),
        ("Records Logged", len(all_recs)),
    ]
    tot_p = sum(s["present"] + s["od"] for s in student_stats.values())
    tot_cond = sum(s["total_conducted"] for s in student_stats.values())
    overall_avg = round((tot_p / tot_cond * 100.0), 1) if tot_cond > 0 else 100.0
    kpis.append(("Overall Attendance", f"{overall_avg}%"))

    col_idx = 2
    for title, val in kpis:
        c_title = ws1.cell(row=7, column=col_idx, value=title)
        c_title.font = Font(name="Segoe UI", size=9, bold=True, color="475569")
        c_title.alignment = Alignment(horizontal="center")
        c_title.fill = header_fill
        c_title.border = thin_border

        c_val = ws1.cell(row=8, column=col_idx, value=val)
        c_val.font = Font(name="Segoe UI", size=13, bold=True, color="1E3A8A")
        c_val.alignment = Alignment(horizontal="center")
        c_val.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        c_val.border = thin_border
        col_idx += 2

    # Table Header (Row 10)
    headers1 = ["Roll No", "Student Name", "Department", "Year", "Periods Taught", "Present", "Absent", "OD", "Attendance %", "Status"]
    for i, h in enumerate(headers1, start=2):
        cell = ws1.cell(row=10, column=i, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center" if i > 3 else "left", vertical="center")
        cell.border = thin_border
    ws1.row_dimensions[10].height = 24

    # Table Rows
    curr_row = 11
    sorted_students = sorted(student_stats.values(), key=lambda x: (x["department"], x["roll_no"]))

    for st in sorted_students:
        pct = round(((st["present"] + st["od"]) / st["total_conducted"]) * 100.0, 1) if st["total_conducted"] > 0 else 100.0
        status_text = "Satisfactory" if pct >= 75.0 else "Shortage (<75%)"

        row_vals = [
            st["roll_no"],
            st["name"],
            st["department"],
            f"Year {st['year']}",
            st["total_conducted"],
            st["present"],
            st["absent"],
            st["od"],
            f"{pct}%",
            status_text
        ]
        for i, val in enumerate(row_vals, start=2):
            cell = ws1.cell(row=curr_row, column=i, value=val)
            cell.font = Font(name="Segoe UI", size=9, bold=(i == 2 or i == 10))
            cell.alignment = Alignment(horizontal="center" if i > 3 else "left", vertical="center")
            cell.border = thin_border

            # Highlight percentage and status
            if i == 10:
                cell.font = Font(name="Segoe UI", size=9, bold=True, color="065F46" if pct >= 75.0 else "991B1B")
                cell.fill = green_fill if pct >= 75.0 else red_fill
            elif i == 11:
                cell.font = Font(name="Segoe UI", size=9, bold=True, color="065F46" if pct >= 75.0 else "991B1B")
                cell.fill = green_fill if pct >= 75.0 else red_fill
        curr_row += 1

    # Auto-fit Sheet 1 columns
    for col in ws1.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # ----------------------------------------------------
    # SHEET 2: DETAILED PERIOD CONDUCTED LOG
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="Detailed Periods Log")
    ws2.views.sheetView[0].showGridLines = True

    ws2.cell(row=2, column=2, value=f"DETAILED PERIOD ATTENDANCE LOG — TAKEN BY {staff_name.upper()} ({staff_initials})").font = Font(name="Segoe UI", size=13, bold=True, color="1E3A8A")
    ws2.cell(row=3, column=2, value=f"Scope: {dept_label} | Timeframe: {date_label} | Total Entries: {len(all_recs)}").font = Font(name="Segoe UI", size=9, italic=True, color="475569")

    headers2 = ["Date", "Day", "Period", "Department", "Year", "Roll No", "Student Name", "Marked Status", "Faculty Name", "Faculty Initials"]
    for i, h in enumerate(headers2, start=2):
        cell = ws2.cell(row=5, column=i, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center" if i in [3, 4, 6, 9, 11] else "left", vertical="center")
        cell.border = thin_border
    ws2.row_dimensions[5].height = 22

    log_row = 6
    amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    for rec, st, dept in all_recs:
        day_str = rec.date.strftime("%a") if rec.date else ""
        row_vals = [
            rec.date.isoformat() if rec.date else "",
            day_str,
            f"Period {rec.hour_number}",
            dept.name,
            f"Year {st.year}",
            st.roll_no,
            st.name,
            rec.status.value.title() if hasattr(rec.status, "value") else str(rec.status).title(),
            staff_name,
            staff_initials
        ]
        for i, val in enumerate(row_vals, start=2):
            cell = ws2.cell(row=log_row, column=i, value=val)
            cell.font = Font(name="Segoe UI", size=9)
            cell.alignment = Alignment(horizontal="center" if i in [3, 4, 6, 9, 11] else "left", vertical="center")
            cell.border = thin_border

            # Status highlight
            if i == 9:
                s_lower = str(val).lower()
                if "present" in s_lower:
                    cell.fill = green_fill
                    cell.font = Font(name="Segoe UI", size=9, bold=True, color="065F46")
                elif "absent" in s_lower:
                    cell.fill = red_fill
                    cell.font = Font(name="Segoe UI", size=9, bold=True, color="991B1B")
                elif "od" in s_lower:
                    cell.fill = amber_fill
                    cell.font = Font(name="Segoe UI", size=9, bold=True, color="92400E")

        log_row += 1

    # Auto-fit Sheet 2 columns
    for col in ws2.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save to BytesIO
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    clean_staff = staff_initials.replace(" ", "_")
    clean_dept = dept_label.replace(" ", "_")
    filename = f"Attendance_Staff_{clean_staff}_{clean_dept}_{filename_suffix}.xlsx"

    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        }
    )


